"""Report maths and rendering. CSV is always available; Excel needs openpyxl, PDF needs reportlab."""
import csv
import io
from datetime import datetime, timezone

from .web import HttpError

PERIOD_FMT = {"hour": "%Y-%m-%d %H:00", "day": "%Y-%m-%d", "month": "%Y-%m", "year": "%Y"}


def period_label(ts_ms, period, tz_min):
    dt = datetime.fromtimestamp((ts_ms + tz_min * 60_000) / 1000.0, tz=timezone.utc)
    return dt.strftime(PERIOD_FMT[period])


def iso(ts_ms):
    if ts_ms is None:
        return ""
    return datetime.fromtimestamp(ts_ms / 1000.0, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def consumption_by_period(db, tag_id, kind, frm, to, period, tz_min):
    """Sum of positive increments (kind='counter') or of drops (kind='level') per period.

    - A decrease of a counter (meter reset / roll-over) is skipped, never counted as negative use.
    - A rise of a fuel level (refuelling) is skipped, never counted as negative use.
    - An increment between two readings is attributed to the period of the later reading.
    - Only Good-quality readings are used; one reading before `frm` is used as the baseline.
    """
    conn = db.conn()
    base = conn.execute("SELECT value FROM samples WHERE tag_id=? AND ts<? AND quality='good' AND value IS NOT NULL "
                        "ORDER BY ts DESC LIMIT 1", (tag_id, int(frm))).fetchone()
    prev = base[0] if base else None
    acc = {}
    for ts, v in conn.execute("SELECT ts, value FROM samples WHERE tag_id=? AND ts>=? AND ts<=? AND quality='good' "
                              "AND value IS NOT NULL ORDER BY ts", (tag_id, int(frm), int(to))):
        label = period_label(ts, period, tz_min)
        acc.setdefault(label, 0.0)
        if prev is not None:
            delta = (v - prev) if kind == "counter" else (prev - v)
            if delta > 0:
                acc[label] += delta
        prev = v
    return acc


def summary_rows(db, tag, frm, to, period, tz_min):
    fmt = PERIOD_FMT[period]
    rows = db.conn().execute(
        "SELECT strftime(?, (ts + ?)/1000, 'unixepoch') AS p, COUNT(*), MIN(value), AVG(value), MAX(value) FROM samples "
        "WHERE tag_id=? AND ts>=? AND ts<=? AND quality='good' AND value IS NOT NULL GROUP BY p ORDER BY p",
        (fmt, tz_min * 60_000, tag["id"], int(frm), int(to))).fetchall()
    cons = None
    if tag["aggregation"] in ("counter", "level"):
        cons = consumption_by_period(db, tag["id"], tag["aggregation"], frm, to, period, tz_min)
    out = []
    for p, n, mn, av, mx in rows:
        out.append({"period": p, "samples": n, "min": mn, "avg": av, "max": mx,
                    "consumption": None if cons is None else cons.get(p, 0.0)})
    return out


def render(fmt, title, columns, rows, meta=None):
    """Return (bytes, content_type, extension)."""
    meta = meta or []
    if fmt == "csv":
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(columns)
        w.writerows(rows)
        return out.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "csv"
    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise HttpError(501, "Excel export needs the 'openpyxl' package on the server (pip install openpyxl)")
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"[:31]
        ws.append(["TBox - Technical Power Systems"])
        ws.append([title])
        for m in meta:
            ws.append(list(m))
        ws.append([])
        hdr_row = ws.max_row + 1
        ws.append(columns)
        for c in ws[hdr_row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2A37")
        for r in rows:
            ws.append(list(r))
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"].font = Font(bold=True)
        ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
        for i, col in enumerate(columns, start=1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)]) + 2
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 60)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    if fmt == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError:
            raise HttpError(501, "PDF export needs the 'reportlab' package on the server (pip install reportlab)")
        if len(rows) > 5000:
            raise HttpError(413, "PDF reports are limited to 5000 rows. Use a shorter range, a longer period, or Excel/CSV.")
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28,
                                title=title, author="TBox")
        st = getSampleStyleSheet()
        story = [Paragraph("<b>TBox</b> - Technical Power Systems", st["Title"]), Paragraph(_esc(title), st["Heading2"])]
        for m in meta:
            story.append(Paragraph("<b>%s</b> %s" % (_esc(m[0]), _esc(m[1] if len(m) > 1 else "")), st["Normal"]))
        story.append(Spacer(1, 10))
        data = [list(map(str, columns))] + [[("" if c is None else (("%.3f" % c).rstrip("0").rstrip(".") if isinstance(c, float) else str(c)))
                                            for c in r] for r in rows]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A37")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F4F8")]),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C5CCD6")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)
        doc.build(story)
        return buf.getvalue(), "application/pdf", "pdf"
    raise HttpError(400, "format must be csv, xlsx or pdf")


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
